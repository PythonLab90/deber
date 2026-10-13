import json
import math
import re
import urllib.error
import urllib.request
from pathlib import Path
from typing import Any

from django.conf import settings


ANALYSIS_FIELDS = (
    "title",
    "topic",
    "executive_summary",
    "key_concepts",
    "important_data",
    "objectives",
    "problems",
    "conclusions",
    "formulas",
    "tables",
    "keywords",
    "study_questions",
)
CHUNK_CHAR_LIMIT = 8_000
_AVAILABLE_GROQ_MODELS: set[str] | None = None
_ACTIVE_GROQ_MODEL: str | None = None
GROQ_USER_AGENT = "groq-python/0.31.0"


class GroqRateLimitError(RuntimeError):
    def __init__(self, message: str, retry_after: int) -> None:
        super().__init__(message)
        self.retry_after = retry_after


def extract_pdf_tables(pdf_path: str | Path) -> list[dict[str, Any]]:
    import pymupdf

    extracted: list[dict[str, Any]] = []
    with pymupdf.open(pdf_path) as document:
        for page_number, page in enumerate(document, start=1):
            page_tables = page.find_tables(strategy="lines").tables
            candidates: list[tuple[list[str], list[list[str]]]] = []
            for table in page_tables:
                rows = [
                    [str(cell).strip() if cell is not None else "" for cell in row]
                    for row in table.extract()
                    if isinstance(row, list)
                ]
                rows = [row for row in rows if any(cell for cell in row)]
                if len(rows) < 2:
                    continue
                headers = rows[0][:30]
                data_rows = [row[:30] for row in rows[1:101]]
                if any(headers) and any(any(cell for cell in row) for row in data_rows):
                    candidates.append((headers, data_rows))

            if not candidates:
                for table in page.find_tables(strategy="text").tables:
                    rows = [
                        [str(cell).strip() if cell is not None else "" for cell in row]
                        for row in table.extract()
                        if isinstance(row, list) and any(cell for cell in row if cell is not None)
                    ]
                    if len(rows) < 2:
                        continue
                    headers = rows[0][:30]
                    data_rows = [row[:30] for row in rows[1:101]]
                    if any(headers) and any(any(cell for cell in row) for row in data_rows):
                        candidates.append((headers, data_rows))

            seen_on_page: set[tuple[tuple[str, ...], ...]] = set()
            for table_number, (headers, data_rows) in enumerate(candidates, start=1):
                signature = tuple(
                    tuple(cell.casefold() for cell in row)
                    for row in [headers, *data_rows]
                )
                if signature in seen_on_page:
                    continue
                seen_on_page.add(signature)
                extracted.append(
                    {
                        "title": f"Tabla {table_number} · página {page_number}",
                        "headers": headers,
                        "rows": data_rows,
                        "source_pages": [page_number],
                    }
                )
    return extracted


def add_detected_pdf_tables(
    analysis: dict[str, Any],
    detected_tables: list[dict[str, Any]],
) -> dict[str, Any]:
    tables = analysis.get("tables")
    if not isinstance(tables, list):
        tables = []
        analysis["tables"] = tables

    existing = {
        tuple(
            tuple(str(cell).strip().casefold() for cell in row)
            for row in [table.get("headers", []), *table.get("rows", [])]
            if isinstance(row, list)
        )
        for table in tables
        if isinstance(table, dict)
        and isinstance(table.get("headers"), list)
        and isinstance(table.get("rows"), list)
    }
    for table in detected_tables:
        signature = tuple(
            tuple(str(cell).strip().casefold() for cell in row)
            for row in [table.get("headers", []), *table.get("rows", [])]
            if isinstance(row, list)
        )
        if signature and signature not in existing:
            tables.append(table)
            existing.add(signature)
    return analysis


def _available_groq_models() -> set[str]:
    global _AVAILABLE_GROQ_MODELS
    if _AVAILABLE_GROQ_MODELS is not None:
        return _AVAILABLE_GROQ_MODELS

    base_url = settings.GROQ_API_URL.rsplit("/", 2)[0]
    request = urllib.request.Request(
        f"{base_url}/models",
        headers={
            "Authorization": f"******",
            "User-Agent": GROQ_USER_AGENT,
            "Accept": "application/json",
        },
        method="GET",
    )
    try:
        with urllib.request.urlopen(request, timeout=settings.GROQ_TIMEOUT_SECONDS) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except (urllib.error.HTTPError, urllib.error.URLError, UnicodeDecodeError, json.JSONDecodeError, OSError):
        return set()

    entries = payload.get("data") if isinstance(payload, dict) else None
    if not isinstance(entries, list):
        return set()
    _AVAILABLE_GROQ_MODELS = {
        item["id"].strip()
        for item in entries
        if isinstance(item, dict)
        and isinstance(item.get("id"), str)
        and item["id"].strip()
    }
    return _AVAILABLE_GROQ_MODELS


def _provider_error(exc: urllib.error.HTTPError) -> tuple[str, str]:
    try:
        body = exc.read(8192).decode("utf-8", errors="replace")
    except OSError:
        body = ""
    try:
        payload = json.loads(body)
    except json.JSONDecodeError:
        payload = {}

    error = payload.get("error", {}) if isinstance(payload, dict) else {}
    if not isinstance(error, dict):
        error = {}
    detail = error.get("message") or error.get("detail") or body
    code = error.get("code", "")
    detail = detail if isinstance(detail, str) else ""
    code = code if isinstance(code, str) else ""
    detail = re.sub(r"(?i)bearer\s+\S+", "Bearer [oculto]", detail)
    detail = re.sub(r"\bgsk_[A-Za-z0-9_-]{10,}\b", "[clave oculta]", detail)
    detail = re.sub(r"\bsk-[A-Za-z0-9_-]{10,}\b", "[clave oculta]", detail)
    return detail[:500].strip(), code[:100].strip()


def _retry_after_seconds(exc: urllib.error.HTTPError, detail: str) -> int:
    retry_after = exc.headers.get("Retry-After") if exc.headers else None
    if retry_after:
        try:
            return min(300, max(1, math.ceil(float(retry_after))))
        except ValueError:
            pass

    match = re.search(
        r"(?:try again in|retry after)\s+(\d+(?:\.\d+)?)\s*(seconds?|secs?|s)\b",
        detail,
        re.IGNORECASE,
    )
    if match:
        return min(300, max(1, math.ceil(float(match.group(1)))))
    return 15


def _parse_json_object(content: Any) -> dict[str, Any] | None:
    if not isinstance(content, str) or not content.strip():
        return None
    cleaned = re.sub(r"^\s*```(?:json)?\s*|\s*```\s*$", "", content, flags=re.IGNORECASE)
    decoder = json.JSONDecoder()
    for index, character in enumerate(cleaned):
        if character != "{":
            continue
        try:
            parsed, _end = decoder.raw_decode(cleaned[index:])
        except json.JSONDecodeError:
            continue
        if isinstance(parsed, dict):
            return parsed
    return None


def _response_content(response_data: Any) -> tuple[Any, str]:
    if not isinstance(response_data, dict):
        return None, ""
    choices = response_data.get("choices")
    if not isinstance(choices, list) or not choices or not isinstance(choices[0], dict):
        return None, ""
    choice = choices[0]
    message = choice.get("message")
    if not isinstance(message, dict):
        return None, str(choice.get("finish_reason") or "")
    content = message.get("content")
    if isinstance(content, list):
        content = "\n".join(
            part.get("text", "")
            for part in content
            if isinstance(part, dict) and isinstance(part.get("text"), str)
        )
    return content, str(choice.get("finish_reason") or "")


def _call_groq(
    prompt: str,
    *,
    model_override: str | None = None,
    allow_fallback: bool = True,
    json_mode: bool = True,
) -> dict[str, Any]:
    global _ACTIVE_GROQ_MODEL
    api_key = settings.GROQ_API_KEY.strip()
    if not api_key:
        raise RuntimeError(
            "No se encontró GROQ_API_KEY. Guárdala en el archivo .env del proyecto y reinicia Django."
        )

    effective_model = model_override or _ACTIVE_GROQ_MODEL or settings.GROQ_MODEL
    request_payload: dict[str, Any] = {
        "model": effective_model,
        "temperature": 0,
        "max_tokens": 2048,
        "messages": [{"role": "user", "content": prompt}],
    }
    if json_mode:
        request_payload["response_format"] = {"type": "json_object"}
    payload = json.dumps(request_payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        settings.GROQ_API_URL,
        data=payload,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": GROQ_USER_AGENT,
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=settings.GROQ_TIMEOUT_SECONDS) as response:
            response_data = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code == 429:
            detail, _provider_code = _provider_error(exc)
            retry_after = _retry_after_seconds(exc, detail)
            message = "Groq alcanzó el límite temporal de solicitudes o tokens."
            if detail:
                message += f" Detalle de Groq: {detail}"
            raise GroqRateLimitError(message, retry_after) from exc
        if exc.code == 400 and json_mode:
            detail, _provider_code = _provider_error(exc)
            if "failed to generate json" in detail.lower():
                retry_prompt = (
                    f"{prompt}\n\n"
                    "REINTENTO: responde únicamente con un objeto JSON válido, sin Markdown, "
                    "sin bloques de código y sin texto antes o después. Cierra todas las "
                    "cadenas, listas y llaves. Si un valor no está disponible, usa una cadena "
                    "vacía o una lista vacía según el esquema."
                )
                try:
                    return _call_groq(
                        retry_prompt,
                        model_override=effective_model,
                        allow_fallback=False,
                        json_mode=False,
                    )
                except RuntimeError as retry_error:
                    if isinstance(retry_error, GroqRateLimitError):
                        raise
                    raise RuntimeError(
                        "Groq no pudo generar JSON en el primer intento ni en el reintento. "
                        f"Detalle del reintento: {retry_error}"
                    ) from exc
        if exc.code == 403 and allow_fallback:
            detail, provider_code = _provider_error(exc)
            model_name = effective_model
            candidates = [
                model
                for model in (
                    "openai/gpt-oss-20b",
                    "qwen/qwen3.8-27b",
                    "openai/gpt-oss-120b",
                )
                if model in _available_groq_models() and model != model_name
            ]
            for fallback_model in candidates:
                try:
                    result = _call_groq(
                        prompt,
                        model_override=fallback_model,
                        allow_fallback=False,
                    )
                    _ACTIVE_GROQ_MODEL = fallback_model
                    return result
                except RuntimeError as fallback_error:
                    detail = f"{fallback_error}. Respuesta del primer modelo: {detail}"
            if candidates:
                message = (
                    "Groq rechazó el modelo configurado y también los modelos alternativos "
                    "que aparecen disponibles para esta clave. Revisa los permisos de inferencia "
                    "del proyecto en Groq Console."
                )
            else:
                message = (
                    "Groq rechazó el acceso y no se encontró otro modelo compatible disponible "
                    "para esta clave. Verifica en Groq Console que la clave esté activa y tenga "
                    "permiso para ejecutar modelos."
                )
            if provider_code:
                message += f" Código de Groq: {provider_code}."
            if detail:
                message += f" Detalle de Groq: {detail}"
            raise RuntimeError(message) from exc

        try:
            error_body = exc.read(8192).decode("utf-8", errors="replace")
            error_data = json.loads(error_body)
        except (OSError, json.JSONDecodeError):
            error_data = {}

        error = error_data.get("error", {}) if isinstance(error_data, dict) else {}
        detail = error.get("message", "") if isinstance(error, dict) else ""
        if not isinstance(detail, str):
            detail = ""
        detail = re.sub(r"(?i)bearer\s+\S+", "Bearer [oculto]", detail)
        detail = re.sub(r"\bgsk_[A-Za-z0-9_-]{10,}\b", "[clave oculta]", detail)
        detail = detail[:500].strip()
        if exc.code == 403:
            explanation = (
                f"Groq denegó el acceso al modelo {effective_model}. "
                "Confirma que la clave pertenezca a una cuenta con acceso a ese modelo."
            )
        elif exc.code == 401:
            explanation = "Groq rechazó la clave. Revisa GROQ_API_KEY en .env y vuelve a iniciar Django."
        else:
            explanation = f"El servicio de análisis respondió con HTTP {exc.code}."
        if detail:
            explanation = f"{explanation} Detalle de Groq: {detail}"
        raise RuntimeError(explanation) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"No se pudo conectar con el servicio de análisis: {exc.reason}") from exc
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeError("El servicio de análisis devolvió una respuesta inválida.") from exc

    content, finish_reason = _response_content(response_data)
    analysis = _parse_json_object(content)
    if analysis is None:
        if json_mode:
            retry_prompt = (
                f"{prompt}\n\n"
                "REINTENTO: responde un único objeto JSON pequeño y válido, sin Markdown "
                "ni texto adicional. Usa listas vacías para categorías sin datos y limita "
                "cada lista a los 5 elementos más importantes. No repitas el documento."
            )
            try:
                return _call_groq(
                    retry_prompt,
                    model_override=effective_model,
                    allow_fallback=False,
                    json_mode=False,
                )
            except RuntimeError as retry_error:
                if isinstance(retry_error, GroqRateLimitError):
                    raise
                raise RuntimeError(
                    "Groq no devolvió un objeto JSON interpretable tras reintentar. "
                    f"Diagnóstico: {retry_error}"
                ) from retry_error
        if finish_reason == "length":
            raise RuntimeError(
                "Groq truncó la respuesta por alcanzar el límite de salida; "
                "el fragmento necesita una respuesta más pequeña."
            )
        if not isinstance(content, str) or not content.strip():
            raise RuntimeError(
                "Groq devolvió una respuesta vacía o sin contenido de texto."
            )
        raise RuntimeError(
            "Groq respondió, pero el contenido no incluye un objeto JSON válido. "
            f"Finalización: {finish_reason or 'no informada'}."
        )
    return analysis


def _make_chunks(pages: list[str]) -> list[list[tuple[int, str]]]:
    chunks: list[list[tuple[int, str]]] = []
    current: list[tuple[int, str]] = []
    current_chars = 0
    for page_number, page_text in enumerate(pages, start=1):
        remaining = page_text.strip()
        while remaining:
            available = CHUNK_CHAR_LIMIT - current_chars
            if available == 0:
                chunks.append(current)
                current = []
                current_chars = 0
                available = CHUNK_CHAR_LIMIT
            piece, remaining = remaining[:available], remaining[available:]
            current.append((page_number, piece))
            current_chars += len(piece)
            if current_chars >= CHUNK_CHAR_LIMIT:
                chunks.append(current)
                current = []
                current_chars = 0
    if current:
        chunks.append(current)
    return chunks


def _chunk_prompt(chunk: list[tuple[int, str]], first_chunk: bool) -> str:
    page_text = "\n\n".join(f"[Página {number}]\n{text}" for number, text in chunk)
    output_template = {
        "title": "string",
        "topic": "string",
        "executive_summary": "string",
        "key_concepts": [{"term": "string", "definition": "string", "source_pages": [1]}],
        "important_data": [{"label": "string", "value": "string", "source_pages": [1]}],
        "objectives": [{"text": "string", "source_pages": [1]}],
        "problems": [{"text": "string", "source_pages": [1]}],
        "conclusions": [{"text": "string", "source_pages": [1]}],
        "formulas": [{"label": "string", "value": "string", "source_pages": [1]}],
        "tables": [
            {
                "title": "string",
                "headers": ["string"],
                "rows": [["string"]],
                "source_pages": [1],
            }
        ],
        "keywords": [{"text": "string", "source_pages": [1]}],
        "study_questions": [
            {"question": "string", "answer": "string", "source_pages": [1]}
        ],
    }
    return (
        "Analiza el fragmento de un documento y devuelve solamente un objeto JSON válido "
        "siguiendo exactamente el esquema indicado. Es un fragmento, no necesariamente el "
        "documento completo: extrae los datos importantes que sí aparecen y no completes "
        "información ausente. No inventes hechos, definiciones, cifras ni respuestas. Si un "
        "dato no está disponible en estas páginas, déjalo como cadena vacía y usa listas vacías "
        "cuando corresponda. Conserva cifras, fechas, nombres y fórmulas tal como aparecen. "
        "Asigna source_pages con los números de las etiquetas [Página N] que respaldan cada "
        "elemento. Para los campos descriptivos title/topic/executive_summary incluye valores "
        "solo en el primer fragmento; el resumen debe describir el contenido visible, no hacer "
        "afirmaciones globales no sustentadas.\n"
        f"Este es el primer fragmento: {json.dumps(first_chunk)}\n"
        f"ESQUEMA: {json.dumps(output_template, ensure_ascii=False)}\n\n"
        f"DOCUMENTO:\n{page_text}"
    )


def _as_text(value: Any, maximum: int = 6000) -> str:
    if not isinstance(value, str):
        return ""
    return value.strip()[:maximum]


def _source_pages(value: Any, page_count: int) -> list[int]:
    if not isinstance(value, list):
        return []
    pages: list[int] = []
    for page in value:
        if isinstance(page, int) and not isinstance(page, bool) and 1 <= page <= page_count:
            if page not in pages:
                pages.append(page)
    return pages


def _items(result: Any, page_count: int, kind: str) -> list[dict[str, Any]]:
    if not isinstance(result, list):
        return []
    output: list[dict[str, Any]] = []
    for item in result:
        if not isinstance(item, dict):
            continue
        if kind == "key_concepts":
            label = _as_text(item.get("term"), 400)
            value = _as_text(item.get("definition"))
        elif kind == "study_questions":
            label = _as_text(item.get("question"), 1000)
            value = _as_text(item.get("answer"))
        elif kind in {"important_data", "formulas"}:
            label = _as_text(item.get("label"), 400)
            value = _as_text(item.get("value"))
        else:
            label = ""
            value = _as_text(item.get("text"))
        if value or label:
            output.append(
                {
                    "label": label,
                    "value": value,
                    "source_pages": _source_pages(item.get("source_pages"), page_count),
                }
            )
    return output


def chunk_count_for_pages(pages: list[str]) -> int:
    return len(_make_chunks(pages))


def analyze_document_chunk(pages: list[str], chunk_index: int) -> dict[str, Any]:
    chunks = _make_chunks(pages)
    if not chunks or chunk_index < 0 or chunk_index >= len(chunks):
        raise ValueError("El documento no contiene texto suficiente para analizar.")

    return _call_groq(_chunk_prompt(chunks[chunk_index], chunk_index == 0))


def merge_document_chunks(analyses: list[dict[str, Any]], page_count: int) -> dict[str, Any]:
    if not analyses:
        raise ValueError("No hay fragmentos analizados para preparar la vista previa.")
    first = analyses[0]
    analysis: dict[str, Any] = {
        "title": _as_text(first.get("title"), 500) or "Título no disponible en el documento",
        "topic": _as_text(first.get("topic"), 1000) or "Tema no disponible en el documento",
        "executive_summary": "\n\n".join(
            summary
            for summary in (_as_text(part.get("executive_summary"), 4000) for part in analyses)
            if summary
        ) or "No se identificó un resumen en el contenido analizado.",
        "key_concepts": [],
        "important_data": [],
        "objectives": [],
        "problems": [],
        "conclusions": [],
        "formulas": [],
        "tables": [],
        "keywords": [],
        "study_questions": [],
    }

    item_fields = (
        "key_concepts",
        "important_data",
        "objectives",
        "problems",
        "conclusions",
        "formulas",
        "keywords",
        "study_questions",
    )
    for part in analyses:
        for field in item_fields:
            analysis[field].extend(_items(part.get(field), page_count, field))
        raw_tables = part.get("tables")
        if isinstance(raw_tables, list):
            for table in raw_tables:
                if not isinstance(table, dict):
                    continue
                headers = table.get("headers")
                rows = table.get("rows")
                if not isinstance(headers, list) or not isinstance(rows, list):
                    continue
                safe_headers = [_as_text(value, 300) for value in headers[:30]]
                safe_rows = [
                    [_as_text(value, 1000) for value in row[:30]]
                    for row in rows[:100]
                    if isinstance(row, list)
                ]
                if safe_headers and safe_rows:
                    analysis["tables"].append(
                        {
                            "title": _as_text(table.get("title"), 300) or "Tabla sin título",
                            "headers": safe_headers,
                            "rows": safe_rows,
                            "source_pages": _source_pages(table.get("source_pages"), page_count),
                        }
                    )

    return analysis


def analyze_document_pages(pages: list[str]) -> dict[str, Any]:
    chunks = _make_chunks(pages)
    analyses = [
        analyze_document_chunk(pages, index)
        for index in range(len(chunks))
    ]
    return merge_document_chunks(analyses, len(pages))


def analysis_to_sections(analysis: dict[str, Any]) -> list[dict[str, Any]]:
    definitions = (
        ("executive_summary", "Resumen ejecutivo", "Resumen del documento", "executive_summary"),
        ("key_concepts", "Conceptos clave", "key_concepts", "key_concepts"),
        ("important_data", "Datos importantes", "important_data", "important_data"),
        ("objectives", "Objetivos", "objectives", "objectives"),
        ("problems", "Problemas", "problems", "problems"),
        ("conclusions", "Conclusiones", "conclusions", "conclusions"),
        ("formulas", "Fórmulas y resultados", "formulas", "formulas"),
        ("tables", "Tablas", "tables", "tables"),
        ("study_questions", "Preguntas de estudio", "study_questions", "study_questions"),
        ("keywords", "Palabras clave", "keywords", "keywords"),
    )
    sections: list[dict[str, Any]] = []
    for key, title, item_key, _ in definitions:
        if key == "executive_summary":
            value = _as_text(analysis.get("executive_summary"), 12000)
            items = [{"label": "", "value": value, "source_pages": []}] if value else []
        elif key == "tables":
            items = []
            for table in analysis.get("tables", []):
                if not isinstance(table, dict):
                    continue
                rows = [" | ".join(table.get("headers", []))]
                rows.extend(" | ".join(row) for row in table.get("rows", []))
                items.append(
                    {
                        "label": _as_text(table.get("title"), 300),
                        "value": "\n".join(rows),
                        "source_pages": table.get("source_pages", []),
                    }
                )
        else:
            items = analysis.get(item_key, [])
        if items:
            sections.append({"key": key, "title": title, "included": True, "items": items})
    return sections


def _clean_payload(payload: Any) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ValueError("No se recibieron los cambios de la vista previa.")
    if payload.get("mode") not in {"intelligent", "complete"}:
        raise ValueError("Selecciona un modo de conversión válido.")
    if payload.get("format") not in {"docx", "pdf", "txt", "md", "xlsx"}:
        raise ValueError("El formato seleccionado no está permitido.")
    if payload.get("detail") not in {"brief", "medium", "detailed"}:
        raise ValueError("Selecciona un nivel de detalle válido.")

    title = _as_text(payload.get("title"), 500)
    topic = _as_text(payload.get("topic"), 1000)
    summary = _as_text(payload.get("executive_summary"), 12000)
    sections = payload.get("sections")
    if not isinstance(sections, list) or len(sections) > 20:
        raise ValueError("La vista previa contiene un número de secciones no válido.")

    safe_sections: list[dict[str, Any]] = []
    seen_keys: set[str] = set()
    for section in sections:
        if not isinstance(section, dict):
            continue
        key = section.get("key")
        is_custom_section = isinstance(key, str) and bool(re.fullmatch(r"custom_[a-z0-9_-]{1,40}", key))
        if (key not in ANALYSIS_FIELDS and not is_custom_section) or key in seen_keys:
            continue
        seen_keys.add(key)
        items = section.get("items")
        if not isinstance(items, list) or len(items) > 200:
            raise ValueError("Una sección contiene demasiados elementos.")
        safe_items = []
        for item in items:
            if not isinstance(item, dict):
                continue
            value = _as_text(item.get("value"), 12000)
            label = _as_text(item.get("label"), 1000)
            if value or label:
                safe_items.append({"label": label, "value": value})
        safe_sections.append(
            {
                "key": key,
                "title": _as_text(section.get("title"), 300),
                "included": section.get("included") is True,
                "items": safe_items,
            }
        )

    return {
        "mode": payload["mode"],
        "format": payload["format"],
        "detail": payload["detail"],
        "title": title,
        "topic": topic,
        "executive_summary": summary,
        "sections": safe_sections,
    }


def _selected_sections(payload: dict[str, Any]) -> list[dict[str, Any]]:
    limit = {"brief": 4, "medium": 10, "detailed": 200}[payload["detail"]]
    output = []
    for section in payload["sections"]:
        if not section["included"]:
            continue
        items = section["items"][:limit]
        if section["key"] == "executive_summary":
            items = [{"label": "", "value": payload["executive_summary"]}]
        if items:
            output.append({**section, "items": items})
    return output


def _report_content(payload: dict[str, Any]) -> str:
    sections = _selected_sections(payload)
    output = [f"# {payload['title']}"]
    if payload["topic"]:
        output.extend(["", f"**Tema principal:** {payload['topic']}"])
    for section in sections:
        output.extend(["", f"## {section['title']}"])
        for item in section["items"]:
            label = item["label"]
            value = item["value"]
            if section["key"] == "tables":
                headers, rows = _parse_table_value(value)
                if headers and rows:
                    if label:
                        output.extend(["", f"### {label}"])
                    output.extend(["", "| " + " | ".join(headers) + " |"])
                    output.append("| " + " | ".join("---" for _ in headers) + " |")
                    output.extend(
                        "| " + " | ".join(row + [""] * max(0, len(headers) - len(row))) + " |"
                        for row in rows
                    )
                    continue
            if label:
                output.extend(["", f"### {label}", "", value])
            else:
                output.extend(["", value])
    return "\n".join(output).strip() + "\n"


def _full_text(pages: list[str]) -> str:
    return "\n\n".join(
        f"Página {number}\n\n{text}" for number, text in enumerate(pages, start=1)
    ).strip()


def _parse_table_value(value: str) -> tuple[list[str], list[list[str]]]:
    lines = [line.strip() for line in value.splitlines() if line.strip()]
    if len(lines) < 2:
        return [], []
    headers = [cell.strip() for cell in lines[0].split("|")]
    rows = [[cell.strip() for cell in line.split("|")] for line in lines[1:]]
    if not headers or not any(headers) or not rows or not any(any(cell for cell in row) for row in rows):
        return [], []
    return headers[:30], [row[:30] for row in rows[:100]]


def _extract_tables(sections: list[dict[str, Any]]) -> list[tuple[str, list[str], list[list[str]]]]:
    output = []
    for section in sections:
        if section["key"] != "tables":
            continue
        for item in section["items"]:
            headers, rows = _parse_table_value(item["value"])
            if headers and rows:
                output.append((item["label"] or "Tabla", headers, rows))
    return output


def _excel_text(value: str) -> str:
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def create_export(payload: dict[str, Any], pages: list[str], original_pdf: Path) -> tuple[bytes, str, str]:
    cleaned = _clean_payload(payload)
    fmt = cleaned["format"]
    mode = cleaned["mode"]
    if fmt == "xlsx" and not _extract_tables(_selected_sections(cleaned)):
        raise ValueError("No hay tablas seleccionadas con datos suficientes para generar Excel.")

    if mode == "complete" and fmt == "pdf":
        return (
            original_pdf.read_bytes(),
            "application/pdf",
            f"{original_pdf.stem}_original.pdf",
        )

    if mode == "complete":
        content = _full_text(pages)
        if not content:
            raise ValueError("El PDF no contiene texto que pueda convertirse.")
        report = f"# {cleaned['title'] or original_pdf.stem}\n\n{content}\n"
        sections = _selected_sections(cleaned) if fmt == "xlsx" else []
    else:
        report = _report_content(cleaned)
        sections = _selected_sections(cleaned)

    if fmt in {"md", "txt"}:
        text = report if fmt == "md" else re.sub(r"(?m)^#{1,6}\s+", "", report).replace("**", "")
        suffix = "md" if fmt == "md" else "txt"
        mime = "text/markdown; charset=utf-8" if fmt == "md" else "text/plain; charset=utf-8"
        return text.encode("utf-8"), mime, f"{original_pdf.stem}.{suffix}"

    if fmt == "xlsx":
        from io import BytesIO

        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        tables = _extract_tables(sections)
        if not tables:
            raise ValueError(
                "No hay tablas seleccionadas con encabezados y filas para generar Excel. "
                "Cambia a conversión inteligente o incluye la sección Tablas."
            )
        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        thin_border = Border(
            left=Side(style="thin", color="D9E2F3"),
            right=Side(style="thin", color="D9E2F3"),
            top=Side(style="thin", color="D9E2F3"),
            bottom=Side(style="thin", color="D9E2F3"),
        )
        header_fill = PatternFill(fill_type="solid", fgColor="0F172A")
        accent_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        soft_fill = PatternFill(fill_type="solid", fgColor="EAF2FF")
        alt_fill = PatternFill(fill_type="solid", fgColor="F8FBFF")

        for table_index, (table_title, headers, rows) in enumerate(tables, start=1):
            sheet = workbook.create_sheet(
                re.sub(r"[\[\]*:/\\?]", " ", table_title)[:31] or f"Tabla {table_index}"
            )
            sheet.sheet_view.showGridLines = False
            sheet.sheet_properties.tabColor = "1F4E78"
            title_text = (table_title or f"Tabla {table_index}").strip()
            title_end_column = min(max(1, len(headers)), 6) if headers else 1
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=title_end_column)
            sheet["A1"] = title_text
            sheet["A1"].font = Font(color="FFFFFF", bold=True, name="Aptos", size=12)
            sheet["A1"].fill = accent_fill
            sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
            sheet["A1"].border = thin_border
            sheet.row_dimensions[1].height = 26

            sheet.freeze_panes = "A3"
            sheet.append([])
            sheet.append([_excel_text(value) for value in headers])
            for row in rows:
                sheet.append([_excel_text(value) for value in row])

            for cell in sheet[3]:
                cell.font = Font(color="FFFFFF", bold=True, name="Aptos")
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                cell.border = thin_border
                sheet.row_dimensions[3].height = 24

            for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
                for cell in row:
                    cell.border = thin_border
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor="F7FAFF" if cell.row % 2 == 0 else "FFFFFF",
                    )
                    if cell.value is not None:
                        cell.font = Font(name="Aptos")

            sheet.auto_filter.ref = sheet.dimensions
            for column_index in range(1, sheet.max_column + 1):
                values = []
                for row_index in range(1, sheet.max_row + 1):
                    value = sheet.cell(row=row_index, column=column_index).value
                    if value is not None:
                        values.append(value)
                if values:
                    width = max(max(len(str(value)) for value in values) + 2, 14)
                    sheet.column_dimensions[get_column_letter(column_index)].width = min(width, 42)

        summary_enabled = bool(cleaned.get("title") or cleaned.get("topic") or cleaned.get("executive_summary"))
        if summary_enabled:
            summary_sheet = workbook.create_sheet("Resumen")
            summary_sheet.sheet_view.showGridLines = False
            summary_sheet.sheet_properties.tabColor = "2E7D8B"
            summary_sheet.merge_cells("A1:F1")
            summary_sheet["A1"] = cleaned["title"] or original_pdf.stem
            summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF", name="Aptos")
            summary_sheet["A1"].fill = accent_fill
            summary_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
            summary_sheet["A1"].border = thin_border
            summary_sheet.row_dimensions[1].height = 26

            row_index = 3
            if cleaned.get("topic"):
                summary_sheet.cell(row=row_index, column=1, value="Tema")
                summary_sheet.cell(row=row_index, column=2, value=cleaned["topic"])
                summary_sheet.cell(row=row_index, column=1).font = Font(bold=True, color="1F1F1F", name="Aptos")
                summary_sheet.cell(row=row_index, column=2).alignment = Alignment(wrap_text=True)
                summary_sheet.cell(row=row_index, column=1).fill = soft_fill
                summary_sheet.cell(row=row_index, column=2).fill = alt_fill
                row_index += 1

            if cleaned.get("executive_summary"):
                summary_sheet.cell(row=row_index, column=1, value="Resumen ejecutivo")
                summary_sheet.cell(row=row_index, column=1).font = Font(bold=True, color="1F1F1F", name="Aptos")
                summary_sheet.cell(row=row_index, column=1).fill = soft_fill
                summary_sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=5)
                summary_sheet.cell(row=row_index, column=1, value=cleaned["executive_summary"])
                summary_sheet.cell(row=row_index, column=1).alignment = Alignment(wrap_text=True)
                summary_sheet.cell(row=row_index, column=1).fill = PatternFill(fill_type="solid", fgColor="EAF2FF")
                row_index += 2

            section_rows: list[tuple[str, str]] = []
            for section in sections:
                if section.get("key") == "tables":
                    continue
                for item in section.get("items", [])[:10]:
                    label = item.get("label") or section.get("title") or "Detalle"
                    value = item.get("value") or ""
                    section_rows.append((str(label), str(value)))

            if section_rows:
                summary_sheet.cell(row=row_index, column=1, value="Detalles")
                summary_sheet.cell(row=row_index, column=1).font = Font(bold=True, color="FFFFFF", name="Aptos")
                summary_sheet.cell(row=row_index, column=1).fill = PatternFill(fill_type="solid", fgColor="5B9BD5")
                summary_sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=2)
                row_index += 1
                for label, value in section_rows:
                    summary_sheet.cell(row=row_index, column=1, value=label)
                    summary_sheet.cell(row=row_index, column=2, value=value)
                    summary_sheet.cell(row=row_index, column=1).font = Font(bold=True, name="Aptos")
                    summary_sheet.cell(row=row_index, column=2).alignment = Alignment(wrap_text=True)
                    row_index += 1

            for column_letter in ("A", "B", "C", "D", "E", "F"):
                values = []
                for row_number in range(1, summary_sheet.max_row + 1):
                    value = summary_sheet.cell(row=row_number, column=ord(column_letter) - 64).value
                    if value is not None:
                        values.append(value)
                if values:
                    summary_sheet.column_dimensions[column_letter].width = min(
                        max(max(len(str(value)) for value in values) + 2, 20), 60
                    )
            summary_sheet.freeze_panes = "A4"

        output = BytesIO()
        workbook.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{original_pdf.stem}_tablas.xlsx",
        )

    if fmt == "docx":
        from io import BytesIO

        from docx import Document
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        document.styles["Normal"].font.name = "Aptos"
        document.styles["Normal"].font.size = Pt(10)
        document.styles["Normal"].font.color.rgb = RGBColor(58, 72, 94)
        document.add_heading(cleaned["title"] or original_pdf.stem, 0)
        if mode == "complete":
            for page_number, page_text in enumerate(pages, start=1):
                document.add_heading(f"Página {page_number}", level=2)
                document.add_paragraph(page_text)
        else:
            if cleaned["topic"]:
                paragraph = document.add_paragraph()
                paragraph.add_run("Tema principal: ").bold = True
                paragraph.add_run(cleaned["topic"])
            for item_section in sections:
                document.add_heading(item_section["title"], level=1)
                for item in item_section["items"]:
                    if item_section["key"] == "tables":
                        headers, rows = _parse_table_value(item["value"])
                        if headers and rows:
                            if item["label"]:
                                document.add_heading(item["label"], level=2)
                            table = document.add_table(rows=1, cols=len(headers))
                            table.style = "Light Shading Accent 1"
                            for cell, value in zip(table.rows[0].cells, headers):
                                cell.text = value
                            for row in rows:
                                cells = table.add_row().cells
                                for index, cell in enumerate(cells):
                                    cell.text = row[index] if index < len(row) else ""
                            continue
                    if item["label"]:
                        document.add_heading(item["label"], level=2)
                    for line in item["value"].splitlines() or [""]:
                        document.add_paragraph(line)
        output = BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{original_pdf.stem}.docx",
        )

    if fmt == "pdf":
        from io import BytesIO

        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        from xml.sax.saxutils import escape

        regular_font = "Helvetica"
        bold_font = "Helvetica-Bold"
        font_dir = Path("C:/Windows/Fonts")
        regular_file = font_dir / "arial.ttf"
        bold_file = font_dir / "arialbd.ttf"
        if regular_file.is_file() and bold_file.is_file():
            pdfmetrics.registerFont(TTFont("DocumentArial", str(regular_file)))
            pdfmetrics.registerFont(TTFont("DocumentArial-Bold", str(bold_file)))
            regular_font = "DocumentArial"
            bold_font = "DocumentArial-Bold"
        stylesheet = getSampleStyleSheet()
        stylesheet.add(
            ParagraphStyle(
                name="ReportTitle",
                parent=stylesheet["Title"],
                fontName=bold_font,
                fontSize=22,
                leading=28,
                textColor=colors.HexColor("#183451"),
                alignment=TA_CENTER,
                spaceAfter=16,
            )
        )
        stylesheet.add(
            ParagraphStyle(
                name="ReportHeading",
                parent=stylesheet["Heading2"],
                fontName=bold_font,
                fontSize=14,
                leading=18,
                textColor=colors.HexColor("#285CA8"),
                spaceBefore=13,
                spaceAfter=6,
            )
        )
        stylesheet.add(
            ParagraphStyle(
                name="ReportBody",
                parent=stylesheet["BodyText"],
                fontName=regular_font,
                fontSize=10,
                leading=15,
                textColor=colors.HexColor("#3A485E"),
                spaceAfter=7,
            )
        )
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=letter,
            rightMargin=0.75 * inch,
            leftMargin=0.75 * inch,
            topMargin=0.75 * inch,
            bottomMargin=0.75 * inch,
            title=cleaned["title"] or original_pdf.stem,
        )
        story = [
            Paragraph(escape(cleaned["title"] or original_pdf.stem), stylesheet["ReportTitle"])
        ]
        if mode == "complete":
            for page_number, page_text in enumerate(pages, start=1):
                story.append(Paragraph(f"Página {page_number}", stylesheet["ReportHeading"]))
                escaped = escape(page_text).replace("\n", "<br/>")
                story.append(Paragraph(escaped or "Página sin texto extraíble.", stylesheet["ReportBody"]))
        else:
            if cleaned["topic"]:
                story.append(Paragraph(f"<b>Tema:</b> {escape(cleaned['topic'])}", stylesheet["ReportBody"]))
            for item_section in sections:
                story.append(Paragraph(escape(item_section["title"]), stylesheet["ReportHeading"]))
                for item in item_section["items"]:
                    if item_section["key"] == "tables":
                        headers, rows = _parse_table_value(item["value"])
                        if headers and rows:
                            if item["label"]:
                                story.append(Paragraph(escape(item["label"]), stylesheet["ReportBody"]))
                            table_data = [
                                [
                                    Paragraph(escape(value), stylesheet["ReportBody"])
                                    for value in row
                                ]
                                for row in [headers, *rows]
                            ]
                            report_table = Table(
                                table_data,
                                repeatRows=1,
                                hAlign="LEFT",
                                splitByRow=1,
                            )
                            report_table.setStyle(
                                TableStyle(
                                    [
                                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#285CA8")),
                                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E1EC")),
                                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                        ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                                        ("TOPPADDING", (0, 0), (-1, -1), 5),
                                        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                                    ]
                                )
                            )
                            story.extend([report_table, Spacer(1, 8)])
                            continue
                    if item["label"]:
                        story.append(Paragraph(escape(item["label"]), stylesheet["ReportHeading"]))
                    for line in item["value"].splitlines() or [""]:
                        story.append(Paragraph(escape(line) or "&nbsp;", stylesheet["ReportBody"]))
        document.build(story)
        return output.getvalue(), "application/pdf", f"{original_pdf.stem}.pdf"

    raise ValueError("No se pudo generar el formato solicitado.")
