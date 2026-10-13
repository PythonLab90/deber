import os
import json
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

import pymupdf
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import RequestFactory, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from .document_workflow import (
    GroqRateLimitError,
    add_detected_pdf_tables,
    create_export,
    extract_pdf_tables,
)
from .schedule_workflow import (
    apply_changes_xlsx,
    build_schedule_prompt,
    call_groq_schedule_analysis,
    extract_xlsx_structure,
)
from .views import extract_pages_from_pdf, load_training_context, parse_extracted_record


class PdfExcelTests(TestCase):
    def test_home_page_loads(self):
        response = self.client.get(reverse("home:index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Convierte tus acuerdos en datos")
        self.assertRegex(response.content.decode(), r"home/styles\.[a-f0-9]+\.css")
        self.assertContains(response, "data-dropzone")
        self.assertContains(response, "Subir PDF")
        self.assertRegex(response.content.decode(), r"home/upload\.[a-f0-9]+\.js")
        self.assertContains(response, "v=20261013-4")
        self.assertRegex(response.content.decode(), r"upload\.[a-f0-9]+\.js\?v=20261013-5")
        self.assertContains(response, "Completa un cronograma Excel")

    def test_upload_analyze_review_workflow(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Acuerdo institucional de prácticas")
        pdf_bytes = document.tobytes()
        document.close()
        analysis = {
            "title": "Acuerdo institucional",
            "topic": "Prácticas preprofesionales",
            "executive_summary": "Acuerdo para organizar las prácticas.",
            "key_concepts": [],
            "important_data": [],
            "objectives": [],
            "problems": [],
            "conclusions": [],
            "formulas": [],
            "tables": [],
            "keywords": [],
            "study_questions": [],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            with override_settings(PDF_JOB_ROOT=Path(temp_dir)):
                response = self.client.post(
                    reverse("home:upload_document"),
                    {
                        "pdf": SimpleUploadedFile(
                            "acuerdo.pdf",
                            pdf_bytes,
                            content_type="application/pdf",
                        )
                    },
                )
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json()["page_count"], 1)
                self.assertEqual(response.json()["chunk_count"], 1)
                self.assertEqual(response.json()["workflow_version"], 2)

                with (
                    patch("home.views.chunk_count_for_pages", return_value=2),
                    patch("home.views.analyze_document_chunk", return_value=analysis),
                ):
                    response = self.client.post(
                        reverse("home:analyze_document"),
                        data=json.dumps({"chunk_index": 0}),
                        content_type="application/json",
                    )
                self.assertEqual(response.status_code, 200)
                self.assertFalse(response.json()["done"])
                self.assertEqual(response.json()["next_chunk"], 1)

                with (
                    patch("home.views.chunk_count_for_pages", return_value=2),
                    patch("home.views.analyze_document_chunk", return_value=analysis),
                ):
                    response = self.client.post(
                        reverse("home:analyze_document"),
                        data=json.dumps({"chunk_index": 1}),
                        content_type="application/json",
                    )
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.json()["done"])

                review_response = self.client.get(reverse("home:review_document"))
                self.assertEqual(review_response.status_code, 200)
                self.assertContains(review_response, "Vista del PDF")
                self.assertContains(review_response, "Vista previa de conversión")
                self.assertContains(review_response, "Acuerdo institucional")

                export_payload = {
                    "mode": "intelligent",
                    "format": "txt",
                    "detail": "medium",
                    "title": "Acuerdo revisado",
                    "topic": "Prácticas",
                    "executive_summary": "Resumen revisado.",
                    "sections": [
                        {
                            "key": "executive_summary",
                            "title": "Resumen ejecutivo",
                            "included": True,
                            "items": [{"label": "", "value": "Resumen original."}],
                        }
                    ],
                }
                response = self.client.post(
                    reverse("home:convert_document"),
                    {"payload": json.dumps(export_payload)},
                )

        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertIn("Acuerdo revisado", response.content.decode("utf-8"))

    @override_settings(GROQ_API_KEY="")
    def test_analyze_endpoint_reports_missing_api_key(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Documento de prueba")
        pdf_bytes = document.tobytes()
        document.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            with override_settings(PDF_JOB_ROOT=Path(temp_dir)):
                upload_response = self.client.post(
                    reverse("home:upload_document"),
                    {
                        "pdf": SimpleUploadedFile(
                            "acuerdo.pdf",
                            pdf_bytes,
                            content_type="application/pdf",
                        )
                    },
                )
                self.assertEqual(upload_response.status_code, 200)
                response = self.client.post(
                    reverse("home:analyze_document"),
                    data=json.dumps({"chunk_index": 0}),
                    content_type="application/json",
                )

        self.assertEqual(response.status_code, 422)
        self.assertIn(".env", response.json()["error"])

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="retired-model")
    def test_groq_403_explains_model_access_and_provider_detail(self):
        from urllib.error import HTTPError

        provider_error = HTTPError(
            "https://api.groq.com/openai/v1/chat/completions",
            403,
            "Forbidden",
            {},
            BytesIO(b'{"error":{"message":"model retired-model is no longer available"}}'),
        )
        with (
            patch("home.document_workflow._AVAILABLE_GROQ_MODELS", None),
            patch("home.document_workflow._ACTIVE_GROQ_MODEL", None),
            patch("home.document_workflow.urllib.request.urlopen", side_effect=provider_error),
        ):
            from .document_workflow import _call_groq

            with self.assertRaisesRegex(RuntimeError, "retired-model") as error:
                _call_groq("test")

        self.assertIn("is no longer available", str(error.exception))
        self.assertNotIn("test-key", str(error.exception))

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="retired-model")
    def test_groq_403_retries_a_model_returned_by_account_catalog(self):
        from urllib.error import HTTPError

        first_error = HTTPError(
            "https://api.groq.com/openai/v1/chat/completions",
            403,
            "Forbidden",
            {},
            BytesIO(b'{"error":{"message":"model access denied"}}'),
        )
        attempted_models = []

        def mock_urlopen(request, timeout):
            self.assertEqual(request.get_header("User-agent"), "groq-python/0.31.0")
            if request.get_method() == "GET":
                return BytesIO(
                    b'{"data":[{"id":"retired-model"},{"id":"qwen/qwen3.8-27b"}]}'
                )
            model = json.loads(request.data.decode("utf-8"))["model"]
            attempted_models.append(model)
            if model == "retired-model":
                raise first_error
            return BytesIO(
                b'{"choices":[{"message":{"content":"{\\"title\\":\\"Documento\\"}"}}]}'
            )

        with (
            patch("home.document_workflow._AVAILABLE_GROQ_MODELS", None),
            patch("home.document_workflow._ACTIVE_GROQ_MODEL", None),
            patch("home.document_workflow.urllib.request.urlopen", side_effect=mock_urlopen),
        ):
            from .document_workflow import _call_groq

            result = _call_groq("test")
            second_result = _call_groq("another test")

        self.assertEqual(result, {"title": "Documento"})
        self.assertEqual(second_result, {"title": "Documento"})
        self.assertEqual(
            attempted_models,
            ["retired-model", "qwen/qwen3.8-27b", "qwen/qwen3.8-27b"],
        )

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    def test_groq_retries_json_generation_failure_without_json_mode(self):
        from urllib.error import HTTPError

        requests = []

        def mock_urlopen(request, timeout):
            payload = json.loads(request.data.decode("utf-8"))
            requests.append(payload)
            if len(requests) == 1:
                raise HTTPError(
                    "https://api.groq.com/openai/v1/chat/completions",
                    400,
                    "Bad Request",
                    {},
                    BytesIO(
                        b'{"error":{"message":"Failed to generate JSON. Please adjust your prompt."}}'
                    ),
                )
            return BytesIO(
                b'{"choices":[{"message":{"content":"{\\"title\\":\\"Documento\\"}"}}]}'
            )

        with patch(
            "home.document_workflow.urllib.request.urlopen",
            side_effect=mock_urlopen,
        ):
            from .document_workflow import _call_groq

            result = _call_groq("Devuelve un objeto JSON sobre el documento.")

        self.assertEqual(result, {"title": "Documento"})
        self.assertIn("response_format", requests[0])
        self.assertNotIn("response_format", requests[1])
        self.assertIn("REINTENTO", requests[1]["messages"][0]["content"])
        self.assertEqual(requests[0]["model"], requests[1]["model"])

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    def test_groq_retries_unparseable_json_content_once(self):
        requests = []

        def mock_urlopen(request, timeout):
            payload = json.loads(request.data.decode("utf-8"))
            requests.append(payload)
            content = (
                "no es JSON"
                if len(requests) == 1
                else 'Resultado solicitado:\n```json\n{"title":"Documento"}\n```'
            )
            return BytesIO(
                json.dumps({"choices": [{"message": {"content": content}}]}).encode("utf-8")
            )

        with patch(
            "home.document_workflow.urllib.request.urlopen",
            side_effect=mock_urlopen,
        ):
            from .document_workflow import _call_groq

            result = _call_groq("Devuelve un objeto JSON sobre el documento.")

        self.assertEqual(result, {"title": "Documento"})
        self.assertEqual(len(requests), 2)
        self.assertNotIn("response_format", requests[1])

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    def test_groq_json_parser_accepts_embedded_object_and_text_blocks(self):
        from .document_workflow import _parse_json_object

        self.assertEqual(
            _parse_json_object('Resultado:\n```json\n{"title":"Documento"}\n```'),
            {"title": "Documento"},
        )
        self.assertEqual(
            _parse_json_object('{"title":"A { B"} trailing text'),
            {"title": "A { B"},
        )
        self.assertIsNone(_parse_json_object(""))
        self.assertIsNone(_parse_json_object("respuesta sin estructura"))

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    def test_groq_reports_truncated_response_without_retry_loop(self):
        requests = []

        def mock_urlopen(request, timeout):
            requests.append(json.loads(request.data.decode("utf-8")))
            response = {
                "choices": [
                    {
                        "message": {"content": "incompleto"},
                        "finish_reason": "length",
                    }
                ]
            }
            return BytesIO(json.dumps(response).encode("utf-8"))

        with patch(
            "home.document_workflow.urllib.request.urlopen",
            side_effect=mock_urlopen,
        ):
            from .document_workflow import _call_groq

            with self.assertRaisesRegex(RuntimeError, "truncó la respuesta"):
                _call_groq("Devuelve un objeto JSON válido.")

        self.assertEqual(len(requests), 2)
        self.assertEqual(requests[0]["max_tokens"], 2048)

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    def test_groq_rate_limit_exposes_retry_after_from_provider(self):
        from urllib.error import HTTPError

        provider_error = HTTPError(
            "https://api.groq.com/openai/v1/chat/completions",
            429,
            "Too Many Requests",
            {"Retry-After": "9"},
            BytesIO(
                b'{"error":{"message":"Rate limit reached. Please try again in 13.4s."}}'
            ),
        )
        body_only_error = HTTPError(
            "https://api.groq.com/openai/v1/chat/completions",
            429,
            "Too Many Requests",
            {},
            BytesIO(
                b'{"error":{"message":"Rate limit reached. Please try again in 13.4025s."}}'
            ),
        )
        with patch(
            "home.document_workflow.urllib.request.urlopen",
            side_effect=[provider_error, body_only_error],
        ):
            from .document_workflow import _call_groq

            with self.assertRaises(GroqRateLimitError) as error:
                _call_groq("test")
            with self.assertRaises(GroqRateLimitError) as body_error:
                _call_groq("test")

        self.assertEqual(error.exception.retry_after, 9)
        self.assertEqual(body_error.exception.retry_after, 14)
        self.assertIn("Rate limit reached", str(error.exception))

    @override_settings(GROQ_API_KEY="test-key")
    def test_analyze_endpoint_returns_retryable_rate_limit_without_losing_chunk(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Documento de prueba con datos.")
        pdf_bytes = document.tobytes()
        document.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            with override_settings(PDF_JOB_ROOT=Path(temp_dir)):
                upload_response = self.client.post(
                    reverse("home:upload_document"),
                    {
                        "pdf": SimpleUploadedFile(
                            "fuente.pdf",
                            pdf_bytes,
                            content_type="application/pdf",
                        )
                    },
                )
                self.assertEqual(upload_response.status_code, 200)
                job_directory = Path(temp_dir) / self.client.session["pdf_job_id"]

                with patch(
                    "home.views.analyze_document_chunk",
                    side_effect=GroqRateLimitError("Límite temporal de Groq.", 14),
                ):
                    response = self.client.post(
                        reverse("home:analyze_document"),
                        data=json.dumps({"chunk_index": 0}),
                        content_type="application/json",
                    )

                self.assertEqual(response.status_code, 429)
                self.assertEqual(response.json()["retry_after"], 14)
                self.assertFalse((job_directory / "analysis_chunk_0000.json").exists())
                self.assertFalse(self.client.session.get("pdf_analyzed", False))

    def test_document_export_formats_and_complete_pdf(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Texto original de la página")
        pdf_bytes = document.tobytes()
        document.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "original.pdf"
            pdf_path.write_bytes(pdf_bytes)
            base_payload = {
                "mode": "intelligent",
                "format": "txt",
                "detail": "detailed",
                "title": "Reporte de prácticas",
                "topic": "Seguimiento académico",
                "executive_summary": "Resumen ajustado por el usuario.",
                "sections": [
                    {
                        "key": "custom_notes",
                        "title": "Conclusiones editadas",
                        "included": True,
                        "items": [{"label": "Resultado", "value": "El acuerdo fue aprobado."}],
                    },
                    {
                        "key": "tables",
                        "title": "Tablas",
                        "included": True,
                        "items": [{"label": "Participantes", "value": "Nombre | Cantidad\nAna | 12"}],
                    },
                ],
            }
            formats = {
                "txt": "Reporte de prácticas",
                "md": "# Reporte de prácticas",
                "docx": b"PK",
                "pdf": b"%PDF",
                "xlsx": b"PK",
            }
            for output_format, expected in formats.items():
                with self.subTest(format=output_format):
                    payload = {**base_payload, "format": output_format}
                    content, _content_type, _filename = create_export(
                        payload,
                        ["Texto original de la página"],
                        pdf_path,
                    )
                    if isinstance(expected, str):
                        self.assertIn(expected, content.decode("utf-8"))
                    else:
                        self.assertTrue(content.startswith(expected))

            complete_payload = {**base_payload, "mode": "complete", "format": "pdf"}
            complete, content_type, _filename = create_export(
                complete_payload,
                ["Texto original de la página"],
                pdf_path,
            )
            self.assertEqual(content_type, "application/pdf")
            self.assertEqual(complete, pdf_bytes)

            complete_xlsx_payload = {
                **base_payload,
                "mode": "complete",
                "format": "xlsx",
            }
            complete_xlsx, xlsx_content_type, _filename = create_export(
                complete_xlsx_payload,
                ["Texto original de la página"],
                pdf_path,
            )
            self.assertEqual(
                xlsx_content_type,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            complete_workbook = load_workbook(BytesIO(complete_xlsx), read_only=True, data_only=True)
            self.assertEqual(list(complete_workbook.sheetnames), ["Participantes", "Resumen"])
            self.assertEqual(
                list(complete_workbook.active.values),
                [("Participantes", None), (None, None), ("Nombre", "Cantidad"), ("Ana", "12")],
            )
            self.assertEqual(complete_workbook["Resumen"]["A1"].value, "Reporte de prácticas")
            self.assertEqual(complete_workbook["Resumen"]["B3"].value, "Seguimiento académico")
            complete_workbook.close()

    def test_xlsx_export_without_usable_selected_tables_returns_validation_error(self):
        payload = {
            "mode": "complete",
            "format": "xlsx",
            "detail": "detailed",
            "title": "Documento",
            "topic": "",
            "executive_summary": "",
            "sections": [
                {
                    "key": "tables",
                    "title": "Tablas",
                    "included": False,
                    "items": [{"label": "Tabla", "value": "Encabezado"}],
                }
            ],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "original.pdf"
            pdf_path.write_bytes(b"%PDF-placeholder")
            with self.assertRaisesRegex(ValueError, "No hay tablas seleccionadas"):
                create_export(payload, ["Texto"], pdf_path)

    def test_convert_endpoint_downloads_xlsx_in_complete_mode_without_server_error(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Texto completo del documento.")
        pdf_bytes = document.tobytes()
        document.close()
        job_id = "a" * 32
        payload = {
            "mode": "complete",
            "format": "xlsx",
            "detail": "detailed",
            "title": "Documento",
            "topic": "",
            "executive_summary": "",
            "sections": [
                {
                    "key": "tables",
                    "title": "Tablas",
                    "included": True,
                    "items": [{"label": "Inscritos", "value": "Programa | Total\nPrácticas | 18"}],
                }
            ],
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            job_directory = Path(temp_dir) / job_id
            job_directory.mkdir()
            (job_directory / "pages.json").write_text(
                json.dumps(["Texto completo del documento."]),
                encoding="utf-8",
            )
            (job_directory / "original.pdf").write_bytes(pdf_bytes)
            session = self.client.session
            session["pdf_job_id"] = job_id
            session["pdf_analyzed"] = True
            session.save()

            with override_settings(PDF_JOB_ROOT=Path(temp_dir)):
                response = self.client.post(
                    reverse("home:convert_document"),
                    {"payload": json.dumps(payload)},
                )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(
            list(workbook.active.values),
            [("Inscritos", None), (None, None), ("Programa", "Total"), ("Prácticas", "18")],
        )
        workbook.close()

    def test_pdf_table_extraction_detects_bordered_and_borderless_tables(self):
        document = pymupdf.open()
        page = document.new_page()
        for x in (50, 250, 450):
            page.draw_line((x, 80), (x, 180))
        for y in (80, 130, 180):
            page.draw_line((50, y), (450, y))
        for position, text in (
            ((60, 110), "Nombre"),
            ((260, 110), "Cantidad"),
            ((60, 160), "Ana"),
            ((260, 160), "12"),
        ):
            page.insert_text(position, text)
        page = document.new_page()
        for position, text in (
            ((60, 100), "Mes"),
            ((260, 100), "Ventas"),
            ((60, 130), "Enero"),
            ((260, 130), "25"),
            ((60, 160), "Febrero"),
            ((260, 160), "31"),
        ):
            page.insert_text(position, text)
        pdf_bytes = document.tobytes()
        document.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "tablas.pdf"
            pdf_path.write_bytes(pdf_bytes)
            tables = extract_pdf_tables(pdf_path)

        self.assertEqual(len(tables), 2)
        self.assertEqual(tables[0]["headers"], ["Nombre", "Cantidad"])
        self.assertEqual(tables[0]["rows"], [["Ana", "12"]])
        self.assertEqual(tables[0]["source_pages"], [1])
        self.assertEqual(tables[1]["headers"], ["Mes", "Ventas"])
        self.assertEqual(tables[1]["rows"], [["Enero", "25"], ["Febrero", "31"]])

    def test_detected_pdf_tables_are_added_to_analysis_without_duplicates(self):
        detected = {
            "title": "Tabla 1 · página 1",
            "headers": ["Nombre", "Cantidad"],
            "rows": [["Ana", "12"]],
            "source_pages": [1],
        }
        analysis = {
            "tables": [
                {
                    "title": "Participantes",
                    "headers": ["Nombre", "Cantidad"],
                    "rows": [["Ana", "12"]],
                    "source_pages": [1],
                }
            ]
        }

        merged = add_detected_pdf_tables(analysis, [detected])

        self.assertEqual(len(merged["tables"]), 1)

    def test_schedule_copy_preserves_template_while_writing_only_approved_empty_cells(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, PatternFill, Side

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Plan"
        worksheet["A1"] = "Cronograma institucional"
        worksheet.merge_cells("A1:B1")
        worksheet["A2"] = "Actividad"
        worksheet["B2"] = "Fecha"
        worksheet["A3"] = "Actividad ya existente"
        worksheet["B3"] = None
        worksheet["B3"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
        worksheet["B3"].border = Border(bottom=Side(style="medium", color="445566"))
        worksheet["B3"].alignment = Alignment(horizontal="center", wrap_text=True)
        worksheet["C3"] = "=1+1"
        worksheet.column_dimensions["B"].width = 31
        worksheet.row_dimensions[3].height = 27
        worksheet.page_setup.orientation = "landscape"
        workbook.create_sheet("Notas")["A1"] = "No modificar"

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "plantilla.xlsx"
            workbook.save(template_path)
            workbook.close()
            original_bytes = template_path.read_bytes()
            structure = extract_xlsx_structure(template_path)
            available = {(item["sheet_name"], item["ref"]) for item in structure["empty_refs"]}
            self.assertIn(("Plan", "B3"), available)
            self.assertNotIn(("Plan", "C3"), available)
            self.assertNotIn(("Plan", "B1"), available)

            completed = apply_changes_xlsx(
                template_path,
                [
                    {
                        "sheet_name": "Plan",
                        "ref": "B3",
                        "proposed_value": "15/11/2026",
                    },
                    {
                        "sheet_name": "Plan",
                        "ref": "A3",
                        "proposed_value": "No debe sobrescribirse",
                    },
                ],
            )

            self.assertEqual(template_path.read_bytes(), original_bytes)
            with zipfile.ZipFile(BytesIO(original_bytes)) as original_archive:
                with zipfile.ZipFile(BytesIO(completed)) as completed_archive:
                    changed_parts = [
                        name
                        for name in original_archive.namelist()
                        if original_archive.read(name) != completed_archive.read(name)
                    ]
            self.assertEqual(changed_parts, ["xl/worksheets/sheet1.xml"])
            result = load_workbook(BytesIO(completed), data_only=False)
            filled = result["Plan"]
            self.assertEqual(filled["A3"].value, "Actividad ya existente")
            self.assertEqual(filled["B3"].value, "15/11/2026")
            self.assertEqual(filled["C3"].value, "=1+1")
            self.assertEqual(filled["B3"].fill.fgColor.rgb, "00FFF2CC")
            self.assertEqual(filled["B3"].border.bottom.style, "medium")
            self.assertEqual(filled["B3"].alignment.horizontal, "center")
            self.assertEqual(filled.column_dimensions["B"].width, 31)
            self.assertEqual(filled.row_dimensions[3].height, 27)
            self.assertIn("A1:B1", {str(cell_range) for cell_range in filled.merged_cells.ranges})
            self.assertEqual(filled.page_setup.orientation, "landscape")
            self.assertEqual(result.sheetnames, ["Plan", "Notas"])
            self.assertEqual(result["Notas"]["A1"].value, "No modificar")
            result.close()

    def test_schedule_structure_prioritizes_contextual_blank_cells_in_large_workbooks(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill

        workbook = Workbook()
        worksheet = workbook.active
        worksheet["AN80"] = "Fin de plantilla"
        worksheet["A70"] = "Objetivo"
        worksheet["B70"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")

        with tempfile.TemporaryDirectory() as temp_dir:
            template_path = Path(temp_dir) / "plantilla.xlsx"
            workbook.save(template_path)
            workbook.close()
            structure = extract_xlsx_structure(template_path)

        priority_refs = structure["empty_refs"][:180]
        objective_cell = next(
            (cell for cell in priority_refs if cell["ref"] == "B70"),
            None,
        )
        self.assertIsNotNone(objective_cell)
        self.assertIn("A70: Objetivo", objective_cell["context"])
        prompt = build_schedule_prompt("Objetivo documentado", structure)
        self.assertIn("B70 (fila: A70: Objetivo", prompt)

    def test_schedule_template_preview_returns_read_only_sheet_cells(self):
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Cronograma"
        worksheet["A1"] = "Actividad"
        worksheet["B1"] = "Fecha"
        worksheet["A2"] = "Revisión"
        template_stream = BytesIO()
        workbook.save(template_stream)
        workbook.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            with override_settings(PDF_JOB_ROOT=Path(temp_dir)):
                response = self.client.post(
                    reverse("home:preview_schedule_template"),
                    {
                        "template": SimpleUploadedFile(
                            "cronograma.xlsx",
                            template_stream.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        )
                    },
                )

        self.assertEqual(response.status_code, 200, response.content)
        preview = response.json()
        self.assertEqual(preview["filename"], "cronograma.xlsx")
        self.assertEqual(preview["sheets"][0]["name"], "Cronograma")
        self.assertEqual(preview["sheets"][0]["rows"][0][0]["value"], "Actividad")
        self.assertEqual(preview["sheets"][0]["rows"][1][1]["value"], "")

    @override_settings(GROQ_API_KEY="test-key", GROQ_MODEL="test-model")
    def test_schedule_analysis_413_reports_clear_user_friendly_error(self):
        from urllib.error import HTTPError

        provider_error = HTTPError(
            "https://api.groq.com/openai/v1/chat/completions",
            413,
            "Request Entity Too Large",
            {},
            BytesIO(b'{"error":{"message":"request body too large"}}'),
        )

        with patch(
            "home.schedule_workflow.urllib.request.urlopen",
            side_effect=provider_error,
        ):
            with self.assertRaisesRegex(RuntimeError, "demasiado grandes|versión más corta"):
                call_groq_schedule_analysis("prompt de prueba")

    @override_settings(GROQ_API_KEY="test-key")
    def test_schedule_upload_preview_and_download_use_original_excel_template(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill

        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Actividad: revisión de expedientes. Fecha: noviembre.")
        pdf_bytes = document.tobytes()
        document.close()

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Cronograma"
        worksheet["A1"] = "Actividad"
        worksheet["B1"] = "Mes"
        worksheet["A2"] = "Revisión"
        worksheet["B2"].fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
        template_stream = BytesIO()
        workbook.save(template_stream)
        workbook.close()
        template_bytes = template_stream.getvalue()

        with tempfile.TemporaryDirectory() as temp_dir:
            with override_settings(PDF_JOB_ROOT=Path(temp_dir)):
                page_response = self.client.get(reverse("home:schedule_page"))
                self.assertEqual(page_response.status_code, 200)
                self.assertContains(page_response, "Completa tu Excel sin cambiar su diseño")
                self.assertContains(page_response, "data-template-preview")
                with patch(
                    "home.views.extract_pages_from_pdf",
                    return_value=["Actividad: revisión de expedientes. Fecha: noviembre."],
                ), patch(
                    "home.views.call_groq_schedule_analysis",
                    return_value={
                        "summary": "Se encontró una actividad con su mes.",
                        "proposed_changes": [
                            {
                                "sheet_name": "Cronograma",
                                "ref": "B2",
                                "proposed_value": "Noviembre",
                                "source_pages": [1],
                                "confidence": 0.95,
                                "reason": "El PDF menciona noviembre.",
                            }
                        ],
                    },
                ):
                    response = self.client.post(
                        reverse("home:analyze_schedule"),
                        {
                            "pdf": SimpleUploadedFile("fuente.pdf", pdf_bytes, content_type="application/pdf"),
                            "template": SimpleUploadedFile(
                                "cronograma.xlsx",
                                template_bytes,
                                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            ),
                        },
                    )
                self.assertEqual(response.status_code, 200, response.content)
                self.assertEqual(response.json()["proposals"][0]["ref"], "B2")

                download = self.client.post(
                    reverse("home:download_completed_schedule"),
                    data=json.dumps(
                        {
                            "changes": [
                                {
                                    "sheet_name": "Cronograma",
                                    "ref": "B2",
                                    "proposed_value": "Noviembre",
                                }
                            ]
                        }
                    ),
                    content_type="application/json",
                )
                self.assertEqual(download.status_code, 200, download.content)
                completed = load_workbook(BytesIO(download.content))
                self.assertEqual(completed["Cronograma"]["B2"].value, "Noviembre")
                self.assertEqual(completed["Cronograma"]["B2"].fill.fgColor.rgb, "00DDEBF7")
                completed.close()
                self.assertEqual(
                    (Path(temp_dir) / self.client.session["schedule_job_id"] / "original.xlsx").read_bytes(),
                    template_bytes,
                )

    def test_parse_extracted_record_accepts_json_fences(self):
        record = parse_extracted_record('```json\n{"Entidad": "UNEMI"}\n```', ["Entidad", "Fecha"])
        self.assertEqual(record, {"Entidad": "UNEMI", "Fecha": ""})

    def test_parse_extracted_record_rejects_invalid_json(self):
        with self.assertRaises(ValueError):
            parse_extracted_record("no es JSON", ["Entidad"])

    def test_training_context_loads_workbook_headers(self):
        headers = load_training_context()["headers"]
        self.assertGreaterEqual(len(headers), 80)
        self.assertEqual(headers[0], "NOMBRE COMERCIAL O SIGLAS ENTIDAD")

    def test_extract_pages_from_text_pdf(self):
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "Acuerdo de prácticas")
        pdf_bytes = document.tobytes()
        document.close()

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "acuerdo.pdf"
            pdf_path.write_bytes(pdf_bytes)
            self.assertEqual(extract_pages_from_pdf(pdf_path), ["Acuerdo de prácticas"])

    def test_extract_pages_rejects_invalid_pdf(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "invalid.pdf"
            pdf_path.write_text("not a PDF", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "no es un PDF válido"):
                extract_pages_from_pdf(pdf_path)

    def test_download_requires_processed_record(self):
        response = self.client.get(reverse("home:download_excel"))
        self.assertEqual(response.status_code, 400)

    @override_settings(GROQ_API_KEY="")
    def test_upload_reports_missing_groq_key(self):
        with patch("home.views.extract_pages_from_pdf", return_value=["Texto del acuerdo"]):
            response = self.client.post(
                reverse("home:extract_pdf"),
                {
                    "pdf": SimpleUploadedFile(
                        "acuerdo.pdf",
                        b"%PDF-1.4 archivo de prueba",
                        content_type="application/pdf",
                    )
                },
            )

        self.assertEqual(response.status_code, 422)
        self.assertContains(response, "GROQ_API_KEY", status_code=422)

    def test_download_excel_contains_extracted_row(self):
        request = RequestFactory().get(reverse("home:download_excel"))
        request.session = {"last_extracted_record": {"Campo A": "Valor"}}
        with override_settings(TRAINING_EXCEL_PATH=self._create_workbook()):
            from .views import download_excel

            response = download_excel(request)

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        rows = list(workbook.active.values)
        self.assertEqual(rows, [("Campo A", "Campo B"), ("Valor", None)])
        workbook.close()

    def _create_workbook(self):
        from tempfile import NamedTemporaryFile

        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["Campo A", "Campo B"])
        temp_file = NamedTemporaryFile(suffix=".xlsx", delete=False)
        workbook.save(temp_file.name)
        temp_file.close()
        self.addCleanup(os.unlink, temp_file.name)
        return temp_file.name
