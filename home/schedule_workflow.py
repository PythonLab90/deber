"""
schedule_workflow.py
Backend logic for the "Completar cronograma con IA" feature.
Extracts schedule structure, calls Groq for AI mapping, and applies approved changes.
"""
import json
import posixpath
import re
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.dom import minidom
from xml.etree import ElementTree
import zipfile

from django.conf import settings


# ── Structure extraction ──────────────────────────────────────────────────────


def extract_xlsx_structure(path: Path) -> dict[str, Any]:
    """Read every sheet while preserving formulas and identifying writable cells."""
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(path, data_only=False, keep_links=True)
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise ValueError("No se pudo abrir la plantilla. Verifica que sea un libro .xlsx válido.") from exc
    cells: list[dict[str, Any]] = []
    empty_refs: list[dict[str, str]] = []
    grid_lines: list[str] = []
    sheet_structures: list[dict[str, Any]] = []
    previews: list[dict[str, Any]] = []
    sheet_names = list(wb.sheetnames)
    merged_ranges: list[str] = []
    max_row = 0
    max_col = 0
    for ws in wb.worksheets:
        scan_rows = min(ws.max_row or 0, 80)
        scan_cols = min(ws.max_column or 0, 40)
        max_row = max(max_row, ws.max_row or 0)
        max_col = max(max_col, ws.max_column or 0)
        merged = [str(rng) for rng in ws.merged_cells.ranges]
        merged_ranges.extend(f"{ws.title}!{rng}" for rng in merged)
        merged_cells = {
            cell.coordinate
            for cell_range in ws.merged_cells.ranges
            for row in ws[cell_range.coord]
            for cell in row
        }
        row_values: dict[int, list[str]] = {}
        column_values: dict[int, list[tuple[int, str]]] = {}

        for row in ws.iter_rows(min_row=1, max_row=scan_rows, max_col=scan_cols):
            for cell in row:
                value = cell.value
                if value is not None and str(value).strip():
                    text_value = str(value).strip()
                    row_values.setdefault(cell.row, []).append(
                        f"{cell.coordinate}: {text_value[:100]}"
                    )
                    column_values.setdefault(cell.column, []).append(
                        (cell.row, text_value[:100])
                    )
                    cells.append(
                        {
                            "sheet_name": ws.title,
                            "ref": cell.coordinate,
                            "row": cell.row,
                            "col": cell.column,
                            "value": text_value[:400],
                        }
                    )

        for row in ws.iter_rows(min_row=1, max_row=min(scan_rows, 30), max_col=scan_cols):
            parts = [
                (str(cell.value).strip()[:60] if cell.value is not None else "(vacío)")
                for cell in row
            ]
            grid_lines.append(f"[{ws.title}] " + " | ".join(parts))

        preview_rows = []
        for row_number, row in enumerate(
            ws.iter_rows(
                min_row=1,
                max_row=min(scan_rows, 40),
                max_col=min(scan_cols, 16),
            ),
            start=1,
        ):
            preview_rows.append(
                [
                    {
                        "ref": cell.coordinate,
                        "value": str(cell.value).strip()[:160] if cell.value is not None else "",
                    }
                    for cell in row
                ]
            )
        previews.append({"name": ws.title, "rows": preview_rows})

        sheet_empty_refs: list[tuple[int, int, int, dict[str, str]]] = []
        for row in ws.iter_rows(min_row=1, max_row=scan_rows, max_col=scan_cols):
            for cell in row:
                if (
                    (cell.value is not None and str(cell.value).strip())
                    or cell.coordinate in merged_cells
                ):
                    continue

                nearby_row_values = row_values.get(cell.row, [])
                nearby_column_values = [
                    (row_number, value)
                    for row_number, value in column_values.get(cell.column, [])
                    if row_number < cell.row
                ]
                row_context = " | ".join(nearby_row_values[:5])[:350]
                column_context = " | ".join(
                    f"{value} (fila {row_number})"
                    for row_number, value in nearby_column_values[-3:]
                )[:250]
                has_adjacent_value = any(
                    ws.cell(row=cell.row, column=column).value not in (None, "")
                    for column in (cell.column - 1, cell.column + 1)
                    if 1 <= column <= scan_cols
                ) or any(
                    ws.cell(row=row_number, column=cell.column).value not in (None, "")
                    for row_number in (cell.row - 1, cell.row + 1)
                    if 1 <= row_number <= scan_rows
                )
                score = (
                    int(bool(nearby_row_values)) * 4
                    + int(bool(nearby_column_values)) * 2
                    + int(has_adjacent_value) * 4
                    + int(cell.style_id > 0) * 3
                )
                entry = {
                    "sheet_name": ws.title,
                    "ref": cell.coordinate,
                    "context": "; ".join(
                        part
                        for part in (
                            f"fila: {row_context}" if row_context else "",
                            f"columna: {column_context}" if column_context else "",
                        )
                        if part
                    )[:500],
                }
                sheet_empty_refs.append(
                    (score, cell.row, cell.column, entry)
                )
        sheet_empty_refs.sort(key=lambda item: (-item[0], item[1], item[2]))
        empty_refs.extend(item[3] for item in sheet_empty_refs)
        sheet_structures.append(
            {
                "name": ws.title,
                "max_row": ws.max_row or 0,
                "max_col": ws.max_column or 0,
                "merged_ranges": merged,
            }
        )
    wb.close()
    return {
        "type": "xlsx",
        "sheets": sheet_names,
        "sheet_structures": sheet_structures,
        "max_row": max_row,
        "max_col": max_col,
        "merged_ranges": merged_ranges[:30],
        "cells": cells[:200],
        "empty_refs": empty_refs[:1200],
        "grid_preview": "\n".join(grid_lines)[:6000],
        "preview": previews[:8],
    }


def extract_docx_structure(path: Path) -> dict[str, Any]:
    """Read a Word document and return its table structure."""
    from docx import Document

    doc = Document(path)
    tables_data: list[dict[str, Any]] = []

    for t_idx, table in enumerate(doc.tables):
        cells_data: list[dict[str, Any]] = []
        grid_lines: list[str] = []
        for r_idx, row in enumerate(table.rows):
            row_parts = []
            for c_idx, cell in enumerate(row.cells):
                text = cell.text.strip()[:300]
                cells_data.append({
                    "row": r_idx,
                    "col": c_idx,
                    "text": text,
                    "is_empty": not bool(text),
                })
                row_parts.append(text[:50] if text else "(vacío)")
            grid_lines.append(" | ".join(row_parts))
        tables_data.append({
            "index": t_idx,
            "rows": len(table.rows),
            "cols": len(table.columns),
            "cells": cells_data,
            "grid_preview": "\n".join(grid_lines)[:2000],
        })

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()][:20]
    return {
        "type": "docx",
        "num_tables": len(tables_data),
        "tables": tables_data,
        "paragraphs": paragraphs,
    }


def extract_pdf_schedule_structure(pages: list[str]) -> dict[str, Any]:
    """Describe a PDF-based schedule as plain text."""
    content = "\n\n".join(
        f"[Página {i + 1}]\n{p.strip()}"
        for i, p in enumerate(pages[:6])
        if p.strip()
    )[:6000]
    return {
        "type": "pdf",
        "pages": len(pages),
        "content": content,
        "warning": (
            "Este cronograma es un PDF. La edición directa es limitada; "
            "se intentará reconstruir el contenido con la mayor fidelidad posible."
        ),
    }


# ── AI analysis ──────────────────────────────────────────────────────────────


def build_schedule_prompt(pdf_text: str, schedule: dict[str, Any]) -> str:
    """Build the Groq prompt for schedule-filling analysis."""
    stype = schedule.get("type", "unknown")
    pdf_excerpt = (pdf_text or "")[:6000]

    if stype == "xlsx":
        empty_cell_refs = ", ".join(
            f"{cell['sheet_name']}!{cell['ref']}"
            + (f" ({cell['context']})" if cell.get("context") else "")
            for cell in schedule.get("empty_refs", [])[:180]
        )
        schedule_desc = (
            f"TIPO: Excel (.xlsx)\n"
            f"Hojas y dimensiones: {json.dumps(schedule.get('sheet_structures', []), ensure_ascii=False)[:1500]}\n"
            f"Rangos combinados: {', '.join(schedule.get('merged_ranges', [])[:8]) or 'ninguno'}\n\n"
            f"VISTA DE CUADRÍCULA:\n{(schedule.get('grid_preview', ''))[:1800]}\n\n"
            "CELDAS VACÍAS DISPONIBLES Y SU CONTEXTO (prioriza las que correspondan "
            "a la información fuente; devuelve siempre una referencia exacta):\n"
            f"{empty_cell_refs[:6500]}"
        )
    elif stype == "docx":
        parts = [f"TIPO: Word (.docx)\nTablas encontradas: {schedule['num_tables']}\n"]
        for t in schedule.get("tables", [])[:3]:
            parts.append(
                f"\nTABLA {t['index'] + 1} ({t['rows']} filas × {t['cols']} columnas):\n"
                f"{(t.get('grid_preview', ''))[:1200]}\n"
            )
        schedule_desc = "".join(parts)[:4000]
    else:
        schedule_desc = (
            f"TIPO: PDF\n"
            f"Páginas: {schedule.get('pages', '?')}\n"
            f"Contenido:\n{(schedule.get('content', ''))[:3500]}\n"
            f"ADVERTENCIA: {schedule.get('warning', '')}"
        )

    return (
        "Eres un asistente experto en análisis documental y gestión de cronogramas. "
        "Se te proporcionan dos documentos:\n"
        "1. El contenido extraído de un PDF fuente (contiene información real: actividades, objetivos, "
        "fechas, responsables, recursos, etc.)\n"
        "2. La estructura de un cronograma plantilla que debe completarse.\n\n"
        "Tu tarea es:\n"
        "- Identificar la información relevante del PDF.\n"
        "- Proponer qué contenido debe colocarse en cada celda vacía del cronograma.\n"
        "- Respetar EXACTAMENTE la estructura existente (no crear nuevas filas ni columnas).\n"
        "- En Excel, devolver sheet_name y ref exactos de una celda vacía disponible; nunca proponer celdas combinadas secundarias.\n"
        "- Para cada dato del documento, busca primero una celda vacía cuyo contexto de fila o columna corresponda al dato. "
        "No descartes una plantilla solo porque las celdas están lejos de la parte superior; la lista de celdas vacías está priorizada por contexto.\n"
        "- Nunca modificar encabezados, celdas ocupadas ni fórmulas.\n"
        "- NUNCA inventar información que no esté en el PDF.\n"
        "- Si no hay datos para una celda, usar null.\n"
        "- Indicar las páginas de origen de cada dato.\n"
        "- Devolver ÚNICAMENTE un objeto JSON válido con este formato exacto:\n\n"
        "FORMATO:\n"
        "{\n"
        '  "proposed_changes": [\n'
        "    {\n"
        '      "sheet_name": "Cronograma",\n'
        '      "ref": "B3",\n'
        '      "table_index": 0,\n'
        '      "row": 2,\n'
        '      "col": 1,\n'
        '      "original_value": "",\n'
        '      "proposed_value": "Texto propuesto",\n'
        '      "source_pages": [1, 2],\n'
        '      "confidence": 0.9,\n'
        '      "reason": "Motivo breve"\n'
        "    }\n"
        "  ],\n"
        '  "extracted_items": [\n'
        '    {"category": "actividad", "text": "...", "source_pages": [1]}\n'
        "  ],\n"
        '  "unmapped_items": [\n'
        '    {"text": "...", "reason": "No se encontró celda correspondiente"}\n'
        "  ],\n"
        '  "warnings": ["Advertencia si aplica"],\n'
        '  "summary": "Descripción breve de lo encontrado y propuesto"\n'
        "}\n\n"
        f"ESTRUCTURA DEL CRONOGRAMA:\n{schedule_desc}\n\n"
        f"CONTENIDO DEL PDF FUENTE:\n{pdf_excerpt}"
    )


def call_groq_schedule_analysis(prompt: str) -> dict[str, Any]:
    """Call Groq to analyze and propose schedule changes."""
    api_key = settings.GROQ_API_KEY.strip()
    if not api_key:
        raise RuntimeError(
            "Falta GROQ_API_KEY. Configúrala en el archivo .env."
        )

    payload = json.dumps(
        {
            "model": settings.GROQ_MODEL,
            "temperature": 0,
            "max_tokens": 4096,
            "response_format": {"type": "json_object"},
            "messages": [{"role": "user", "content": prompt}],
        },
        ensure_ascii=False,
    ).encode("utf-8")

    req = urllib.request.Request(
        settings.GROQ_API_URL,
        data=payload,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": "groq-python/0.31.0",
        },
        method="POST",
    )

    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code == 413:
            raise RuntimeError(
                "El PDF o la plantilla son demasiado grandes para el análisis de IA. "
                "Reduce el tamaño del documento fuente o prueba con una versión más corta del archivo."
            ) from exc
        raise RuntimeError(f"Groq respondió con HTTP {exc.code}.") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"No se pudo conectar con Groq: {exc.reason}") from exc
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise RuntimeError("Groq devolvió una respuesta inválida.") from exc

    try:
        content = data["choices"][0]["message"]["content"]
        result = json.loads(content)
    except (KeyError, IndexError, json.JSONDecodeError) as exc:
        raise RuntimeError("La respuesta de Groq no contiene JSON válido.") from exc

    if not isinstance(result, dict):
        raise RuntimeError("La respuesta de Groq no tiene el formato esperado.")
    return result


# ── Change application ────────────────────────────────────────────────────────


def _clean_changes(raw_changes: list[Any]) -> list[dict[str, Any]]:
    """Validate and sanitise the list of approved changes."""
    clean: list[dict[str, Any]] = []
    seen: set[str] = set()
    for item in raw_changes:
        if not isinstance(item, dict):
            continue
        ref = str(item.get("ref", "")).strip().upper()
        sheet_name = str(item.get("sheet_name", "")).strip()
        table_index = item.get("table_index", 0)
        row = item.get("row")
        col = item.get("col")
        value = str(item.get("proposed_value", "")).strip()[:2000]
        if not value:
            continue
        if any(not _is_valid_xml_character(character) for character in value):
            raise ValueError("Una propuesta contiene caracteres que Excel no puede guardar.")
        key = f"{sheet_name}:{table_index}:{row}:{col}:{ref}"
        if key in seen:
            continue
        seen.add(key)
        clean.append({
            "ref": ref,
            "sheet_name": sheet_name,
            "original_value": str(item.get("original_value", ""))[:2000],
            "table_index": int(table_index) if isinstance(table_index, (int, float)) else 0,
            "row": int(row) if isinstance(row, (int, float)) else None,
            "col": int(col) if isinstance(col, (int, float)) else None,
            "proposed_value": value,
        })
    return clean


def apply_changes_xlsx(template_path: Path, raw_changes: list[Any]) -> bytes:
    """Patch only approved blank-cell values in a copy of the original XLSX package."""
    changes = _clean_changes(raw_changes)
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    main_tag = f"{{{spreadsheet_ns}}}"

    with zipfile.ZipFile(template_path, "r") as source:
        workbook_root = ElementTree.fromstring(source.read("xl/workbook.xml"))
        relationships_root = ElementTree.fromstring(source.read("xl/_rels/workbook.xml.rels"))
        relationship_targets = {
            relationship.attrib["Id"]: relationship.attrib["Target"]
            for relationship in relationships_root.findall(f"{{{package_rel_ns}}}Relationship")
        }
        sheet_parts: dict[str, str] = {}
        for sheet in workbook_root.findall(f".//{main_tag}sheet"):
            relation_id = sheet.attrib.get(f"{{{office_rel_ns}}}id")
            target = relationship_targets.get(relation_id or "")
            if not target:
                continue
            part = (
                target.lstrip("/")
                if target.startswith("/")
                else posixpath.normpath(posixpath.join("xl", target))
            )
            sheet_parts[sheet.attrib["name"]] = part

        changed_parts: dict[str, bytes] = {}
        applied = 0
        for sheet_name, sheet_changes in _changes_by_sheet(changes).items():
            part = sheet_parts.get(sheet_name)
            if not part:
                continue
            document = minidom.parseString(source.read(part))
            sheet_data_nodes = document.getElementsByTagNameNS(spreadsheet_ns, "sheetData")
            if not sheet_data_nodes:
                continue
            sheet_data = sheet_data_nodes[0]
            for change in sheet_changes:
                ref = change["ref"]
                if not re.fullmatch(r"[A-Z]{1,3}[0-9]{1,7}", ref):
                    continue
                row_number = int(re.search(r"[0-9]+$", ref).group())
                cell_column = _column_number(re.match(r"[A-Z]+", ref).group())
                row = next(
                    (
                        node for node in sheet_data.childNodes
                        if node.nodeType == node.ELEMENT_NODE
                        and node.localName == "row"
                        and node.getAttribute("r") == str(row_number)
                    ),
                    None,
                )
                if row is None:
                    row = document.createElementNS(spreadsheet_ns, _qualified_tag(sheet_data, "row"))
                    row.setAttribute("r", str(row_number))
                    following_row = next(
                        (
                            node for node in sheet_data.childNodes
                            if node.nodeType == node.ELEMENT_NODE
                            and node.localName == "row"
                            and int(node.getAttribute("r") or "0") > row_number
                        ),
                        None,
                    )
                    sheet_data.insertBefore(row, following_row)

                cell = next(
                    (
                        node for node in row.childNodes
                        if node.nodeType == node.ELEMENT_NODE
                        and node.localName == "c"
                        and node.getAttribute("r") == ref
                    ),
                    None,
                )
                if cell is not None and any(
                    node.nodeType == node.ELEMENT_NODE
                    and node.namespaceURI == spreadsheet_ns
                    and node.localName in {"f", "v", "is"}
                    and (
                        node.localName == "f"
                        or _xml_text_content(node).strip()
                    )
                    for node in cell.childNodes
                ):
                    continue
                if cell is None:
                    cell = document.createElementNS(spreadsheet_ns, _qualified_tag(row, "c"))
                    cell.setAttribute("r", ref)
                    following_cell = next(
                        (
                            node for node in row.childNodes
                            if node.nodeType == node.ELEMENT_NODE
                            and node.localName == "c"
                            and _column_number(re.match(r"[A-Z]+", node.getAttribute("r")).group())
                            > cell_column
                        ),
                        None,
                    )
                    row.insertBefore(cell, following_cell)
                for node in list(cell.childNodes):
                    if (
                        node.nodeType == node.ELEMENT_NODE
                        and node.namespaceURI == spreadsheet_ns
                        and node.localName in {"f", "v", "is"}
                    ):
                        cell.removeChild(node)
                cell.setAttribute("t", "inlineStr")
                inline_string = document.createElementNS(
                    spreadsheet_ns,
                    _qualified_tag(cell, "is"),
                )
                text_node = document.createElementNS(
                    spreadsheet_ns,
                    _qualified_tag(cell, "t"),
                )
                text_node.appendChild(document.createTextNode(change["proposed_value"]))
                inline_string.appendChild(text_node)
                cell.appendChild(inline_string)
                applied += 1
            changed_parts[part] = document.toxml(encoding="utf-8")
            document.unlink()

        if not applied:
            raise ValueError("No se pudo aplicar ningún cambio a una celda vacía de la plantilla.")
        output = BytesIO()
        with zipfile.ZipFile(output, "w") as destination:
            for entry in source.infolist():
                destination.writestr(
                    entry,
                    changed_parts.get(entry.filename, source.read(entry.filename)),
                )
    return output.getvalue()


def _changes_by_sheet(changes: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    by_sheet: dict[str, list[dict[str, Any]]] = {}
    for change in changes:
        sheet_name = change["sheet_name"]
        if sheet_name:
            by_sheet.setdefault(sheet_name, []).append(change)
    return by_sheet


def _qualified_tag(parent: Any, local_name: str) -> str:
    return f"{parent.prefix}:{local_name}" if parent.prefix else local_name


def _xml_text_content(node: Any) -> str:
    return "".join(
        child.data if child.nodeType in {child.TEXT_NODE, child.CDATA_SECTION_NODE}
        else _xml_text_content(child) if child.nodeType == child.ELEMENT_NODE
        else ""
        for child in node.childNodes
    )


def _column_number(column: str) -> int:
    number = 0
    for character in column:
        number = number * 26 + ord(character) - ord("A") + 1
    return number


def _is_valid_xml_character(character: str) -> bool:
    codepoint = ord(character)
    return (
        codepoint in {0x9, 0xA, 0xD}
        or 0x20 <= codepoint <= 0xD7FF
        or 0xE000 <= codepoint <= 0xFFFD
        or 0x10000 <= codepoint <= 0x10FFFF
    )


def apply_changes_docx(template_path: Path, raw_changes: list[Any]) -> bytes:
    """Apply approved changes to a Word document, preserving structure."""
    from docx import Document

    doc = Document(template_path)
    changes = _clean_changes(raw_changes)

    # Build map: (table_idx, row, col) -> value
    cell_map: dict[tuple[int, int, int], str] = {}
    for change in changes:
        t = change["table_index"]
        r = change["row"]
        c = change["col"]
        if r is not None and c is not None:
            cell_map[(t, r, c)] = change["proposed_value"]

    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                key = (t_idx, r_idx, c_idx)
                if key in cell_map:
                    current = cell.text.strip()
                    new_val = cell_map[key]
                    if not current:
                        # Write to first paragraph
                        if cell.paragraphs:
                            cell.paragraphs[0].clear()
                            cell.paragraphs[0].add_run(new_val)
                        else:
                            cell.add_paragraph(new_val)
                    else:
                        # Append with separator if cell already has content
                        cell.add_paragraph(new_val)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()
